"""
Testes da exportação ``?formato=xlsx``.

Abrimos a planilha gerada com o openpyxl e conferimos o que importa para quem
vai usá-la: cabeçalho na linha certa, filtros descritos no topo, valores com o
**tipo nativo** (data/número, não texto), rodapé de totais e autofiltro.
"""
from datetime import date
from io import BytesIO
from numbers import Number

import pytest
from openpyxl import load_workbook

from apps.relatorios.excel import CONTENT_TYPE_XLSX, FORMATO_MOEDA

pytestmark = pytest.mark.django_db

URL_PACIENTES = "/api/relatorios/pacientes/"
URL_AGENDAMENTOS = "/api/relatorios/agendamentos/"
URL_PRODUCAO = "/api/relatorios/producao/"
URL_REPASSE = "/api/relatorios/repasse/"

LINHA_CABECALHO = 4
LINHA_PRIMEIRO_DADO = 5


def abrir_planilha(resposta):
    """Carrega a planilha da resposta HTTP."""
    return load_workbook(BytesIO(resposta.content)).active


def test_download_tem_content_type_e_nome_de_arquivo(api, direcao, agendamentos):
    """A resposta é um anexo .xlsx com a data no nome do arquivo."""
    resposta = api(direcao).get(URL_AGENDAMENTOS, {"formato": "xlsx"})

    assert resposta.status_code == 200
    assert resposta["Content-Type"] == CONTENT_TYPE_XLSX
    assert "attachment;" in resposta["Content-Disposition"]
    assert f"relatorio_agendamentos_{date.today():%Y-%m-%d}.xlsx" in (
        resposta["Content-Disposition"]
    )
    # Sem esse header o navegador não lê o nome do arquivo em requisição CORS.
    assert "Content-Disposition" in resposta["Access-Control-Expose-Headers"]


def test_planilha_tem_titulo_filtros_e_cabecalho(api, direcao, agendamentos):
    """Topo da planilha: título, filtros aplicados e cabeçalho das colunas."""
    resposta = api(direcao).get(
        URL_AGENDAMENTOS,
        {"formato": "xlsx", "data_inicio": "2026-11-01", "data_fim": "2026-11-30"},
    )
    ws = abrir_planilha(resposta)

    assert ws["A1"].value == "Relatório de Agendamentos"
    assert "01/11/2026" in ws["A2"].value and "30/11/2026" in ws["A2"].value
    assert ws.cell(row=LINHA_CABECALHO, column=1).value == "Data"
    assert ws.cell(row=LINHA_CABECALHO, column=5).value == "Paciente"


def test_planilha_escreve_data_como_data_e_nao_texto(api, direcao, agendamentos):
    """
    A coluna de data traz um ``date`` real, com formato brasileiro aplicado.

    É o que permite ordenar e filtrar por data na planilha depois de baixada.
    """
    resposta = api(direcao).get(
        URL_AGENDAMENTOS, {"formato": "xlsx", "status": "ATENDIDO"}
    )
    celula = abrir_planilha(resposta).cell(row=LINHA_PRIMEIRO_DADO, column=1)

    assert celula.value.date() == date(2026, 11, 2)
    assert celula.number_format == "DD/MM/YYYY"


def test_planilha_escreve_moeda_como_numero(api, direcao, agendamentos, pagamento):
    """Valores monetários são numéricos e formatados como moeda."""
    resposta = api(direcao).get(
        URL_PRODUCAO, {"formato": "xlsx", "motivo": "Atendimento Realizado"}
    )
    ws = abrir_planilha(resposta)

    coluna_valor = [
        indice
        for indice in range(1, ws.max_column + 1)
        if ws.cell(row=LINHA_CABECALHO, column=indice).value == "Valor produzido"
    ][0]
    celula = ws.cell(row=LINHA_PRIMEIRO_DADO, column=coluna_valor)

    # Numérico (o Excel não distingue int de float na leitura) e não texto —
    # é o que permite somar a coluna na planilha.
    assert celula.value == 200
    assert isinstance(celula.value, Number)
    assert celula.number_format == FORMATO_MOEDA


def test_planilha_tem_rodape_de_totais(api, direcao, agendamentos, pagamento):
    """A última linha traz a contagem de registros e a soma das colunas somáveis."""
    resposta = api(direcao).get(URL_PRODUCAO, {"formato": "xlsx"})
    ws = abrir_planilha(resposta)

    # 4 linhas de topo/cabeçalho + 2 lançamentos + 1 de totais.
    linha_totais = ws.max_row
    assert ws.cell(row=linha_totais, column=1).value == "2 registro(s)"

    coluna_valor = [
        indice
        for indice in range(1, ws.max_column + 1)
        if ws.cell(row=LINHA_CABECALHO, column=indice).value == "Valor produzido"
    ][0]
    assert ws.cell(row=linha_totais, column=coluna_valor).value == 400.0


def test_planilha_tem_autofiltro_e_painel_congelado(api, direcao, agendamentos):
    """A planilha já chega navegável: autofiltro no cabeçalho e header fixo."""
    resposta = api(direcao).get(URL_AGENDAMENTOS, {"formato": "xlsx"})
    ws = abrir_planilha(resposta)

    assert ws.auto_filter.ref is not None
    assert ws.auto_filter.ref.startswith(f"A{LINHA_CABECALHO}")
    assert ws.freeze_panes == f"A{LINHA_PRIMEIRO_DADO}"


def test_planilha_vazia_nao_quebra(api, direcao):
    """Sem dados, a planilha sai apenas com título e cabeçalho — sem erro."""
    resposta = api(direcao).get(URL_PACIENTES, {"formato": "xlsx"})
    ws = abrir_planilha(resposta)

    assert resposta.status_code == 200
    assert ws["A1"].value == "Relatório de Pacientes"
    assert ws.cell(row=LINHA_CABECALHO, column=1).value == "Nome completo"
    # Sem linhas de dados, também não há rodapé de totais.
    assert ws.cell(row=LINHA_PRIMEIRO_DADO, column=1).value is None


def test_repasse_consolidado_exporta_colunas_agregadas(
    api, direcao, agendamentos, pagamento
):
    """A visão consolidada troca o conjunto de colunas da planilha."""
    resposta = api(direcao).get(URL_REPASSE, {"formato": "xlsx", "agrupar": "true"})
    ws = abrir_planilha(resposta)

    cabecalhos = [
        ws.cell(row=LINHA_CABECALHO, column=indice).value
        for indice in range(1, ws.max_column + 1)
    ]
    assert "Atendimentos pagos" in cabecalhos
    assert "Total a repassar" in cabecalhos
    # Coluna exclusiva da visão detalhada não aparece aqui.
    assert "Paciente" not in cabecalhos
    assert "relatorio_repasse_consolidado" in resposta["Content-Disposition"]


def test_planilha_de_pacientes_traz_o_cadastro(api, direcao, paciente):
    """Sanidade do relatório de pacientes em planilha."""
    resposta = api(direcao).get(URL_PACIENTES, {"formato": "xlsx"})
    ws = abrir_planilha(resposta)

    assert ws.cell(row=LINHA_PRIMEIRO_DADO, column=1).value == "João da Silva"


@pytest.mark.parametrize(
    "url", [URL_PACIENTES, URL_AGENDAMENTOS, URL_PRODUCAO, URL_REPASSE]
)
def test_exportacao_respeita_a_permissao(api, recepcao, financeiro, url):
    """
    O ``?formato=xlsx`` não é um contorno da permissão.

    Quem recebe 403 no JSON também recebe 403 no download.
    """
    from apps.relatorios.tests.test_permissoes import FINANCEIROS

    usuario = recepcao if url in FINANCEIROS else financeiro
    assert api(usuario).get(url, {"formato": "xlsx"}).status_code == 403


def test_datetime_vai_para_a_planilha_no_fuso_da_clinica(
    api, direcao, agendamentos
):
    """
    Regressão de fuso: ``USE_TZ=True``, então ``criado_em`` é UTC no banco.

    Na planilha o valor tem de sair no fuso do projeto (``America/Belem``,
    UTC-3) e **sem** ``tzinfo`` — o Excel não tem tipo "datetime com fuso" e
    mostraria o horário de UTC. No JSON, ao contrário, o offset é preservado
    (é o que ``formatarDataHora`` no frontend espera).
    """
    from django.utils import timezone

    from apps.clinica.models import Agendamento

    criado_em = Agendamento.objects.get(
        pk=agendamentos["atendido"].pk
    ).criado_em
    esperado = timezone.localtime(criado_em)

    resposta = api(direcao).get(
        URL_AGENDAMENTOS, {"formato": "xlsx", "status": "ATENDIDO"}
    )
    ws = abrir_planilha(resposta)
    coluna = [
        indice
        for indice in range(1, ws.max_column + 1)
        if ws.cell(row=LINHA_CABECALHO, column=indice).value == "Agendado em"
    ][0]
    celula = ws.cell(row=LINHA_PRIMEIRO_DADO, column=coluna).value

    assert celula.tzinfo is None
    assert (celula.hour, celula.minute) == (esperado.hour, esperado.minute)


def test_json_preserva_o_offset_do_datetime(api, direcao, agendamentos):
    """Contraparte do teste acima: no JSON o datetime sai com fuso."""
    resposta = api(direcao).get(URL_AGENDAMENTOS, {"status": "ATENDIDO"})
    criado_em = resposta.data["linhas"][0]["criado_em"]

    assert criado_em.tzinfo is not None


def test_texto_com_igual_nao_vira_formula(api, direcao, agendamentos, paciente):
    """
    Segurança: campo livre não pode virar fórmula na planilha (CWE-1236).

    O openpyxl infere fórmula de qualquer string iniciada por ``=``. Uma
    observação digitada como ``=HYPERLINK(...)`` chegaria executável na máquina
    de quem abrisse o arquivo.
    """
    agendamento = agendamentos["atendido"]
    agendamento.observacoes = '=HYPERLINK("http://malicioso.example","clique")'
    agendamento.save()

    resposta = api(direcao).get(
        URL_AGENDAMENTOS, {"formato": "xlsx", "status": "ATENDIDO"}
    )
    ws = abrir_planilha(resposta)
    coluna = [
        indice
        for indice in range(1, ws.max_column + 1)
        if ws.cell(row=LINHA_CABECALHO, column=indice).value == "Observações"
    ][0]
    celula = ws.cell(row=LINHA_PRIMEIRO_DADO, column=coluna)

    # Texto, não fórmula — e o conteúdo original preservado (sem apóstrofo).
    assert celula.data_type == "s"
    assert celula.value == '=HYPERLINK("http://malicioso.example","clique")'
