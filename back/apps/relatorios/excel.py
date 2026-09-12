"""
Geração das planilhas ``.xlsx`` dos relatórios (openpyxl).

Um único gerador serve os quatro relatórios: cada um descreve suas colunas
(``Coluna``) e entrega as linhas como dicionários. Assim o layout da planilha
(cabeçalho, congelamento, autofiltro, formato de moeda/data, rodapé de totais)
fica definido em um lugar só.

O valor vai para a célula com o **tipo nativo** (``date``, ``Decimal``, ``int``)
e a formatação é aplicada via ``number_format`` — nunca como texto pré-formatado.
É o que permite somar, ordenar e filtrar na planilha depois de baixada.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterable, Sequence

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

CONTENT_TYPE_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

#: Formatos numéricos do Excel usados pelas colunas.
FORMATO_MOEDA = 'R$ #,##0.00'
FORMATO_DATA = "DD/MM/YYYY"
FORMATO_DATA_HORA = "DD/MM/YYYY HH:MM"
FORMATO_INTEIRO = "#,##0"

# Paleta do cabeçalho — alinhada ao tema neutro (zinc) da interface.
_COR_CABECALHO = "18181B"
_COR_TEXTO_CABECALHO = "FFFFFF"
_COR_TOTAIS = "F4F4F5"
_COR_BORDA = "D4D4D8"

#: Teto de linhas por exportação.
#:
#: Um relatório sem período (ou com anos de dados) carregaria tudo em memória e
#: montaria uma planilha de dezenas de MB — derrubando o worker. Ao estourar, a
#: API devolve 400 pedindo um recorte menor, em vez de truncar em silêncio: uma
#: planilha cortada sem aviso é pior que um erro, porque vai para o fechamento
#: contábil como se estivesse completa.
MAX_LINHAS_RELATORIO = 20_000

_LINHA_TITULO = 1
_LINHA_PERIODO = 2
_LINHA_CABECALHO = 4


@dataclass(frozen=True)
class Coluna:
    """
    Descreve uma coluna do relatório — usada tanto no JSON quanto no Excel.

    - ``chave``:   nome do campo no dicionário da linha (e na resposta JSON).
    - ``titulo``:  cabeçalho exibido na planilha.
    - ``largura``: largura da coluna em caracteres.
    - ``formato``: ``number_format`` do Excel (moeda, data, inteiro).
    - ``somar``:   inclui a coluna no rodapé de totais.
    """

    chave: str
    titulo: str
    largura: int = 20
    formato: str | None = None
    somar: bool = False


@dataclass(frozen=True)
class Relatorio:
    """Um relatório pronto para virar JSON ou planilha."""

    titulo: str
    nome_arquivo: str
    colunas: Sequence[Coluna]
    linhas: Sequence[dict[str, Any]]
    #: Descrição dos filtros aplicados, exibida no topo da planilha.
    filtros_descricao: str = ""
    #: Totais agregados (ex.: ``{"valor": Decimal("1200.00")}``).
    totais: dict[str, Any] = field(default_factory=dict)

    @property
    def total_registros(self) -> int:
        return len(self.linhas)


def _celula_valor(valor: Any) -> Any:
    """
    Converte o valor para um tipo que o openpyxl escreve nativamente.

    Datetimes chegam com fuso (``USE_TZ=True``, portanto UTC). O Excel não tem
    tipo "datetime com fuso": convertemos para o fuso do projeto
    (``America/Belem``) e removemos o ``tzinfo``, senão a planilha mostraria o
    horário em UTC. O JSON, por outro lado, mantém o offset — é o que o
    frontend espera.
    """
    if isinstance(valor, Decimal):
        return float(valor)
    if isinstance(valor, datetime):
        if timezone.is_naive(valor):
            return valor
        return timezone.localtime(valor).replace(tzinfo=None)
    if isinstance(valor, (date, int, float, str)) or valor is None:
        return valor
    return str(valor)


def _escrever_titulo(ws: Worksheet, relatorio: Relatorio) -> None:
    """Título do relatório e linha de filtros aplicados, mesclados na largura."""
    ultima = get_column_letter(len(relatorio.colunas))

    ws.merge_cells(f"A{_LINHA_TITULO}:{ultima}{_LINHA_TITULO}")
    celula_titulo = ws[f"A{_LINHA_TITULO}"]
    celula_titulo.value = relatorio.titulo
    celula_titulo.font = Font(bold=True, size=14)
    celula_titulo.alignment = Alignment(vertical="center")
    ws.row_dimensions[_LINHA_TITULO].height = 22

    ws.merge_cells(f"A{_LINHA_PERIODO}:{ultima}{_LINHA_PERIODO}")
    celula_filtros = ws[f"A{_LINHA_PERIODO}"]
    celula_filtros.value = relatorio.filtros_descricao or "Sem filtros aplicados"
    celula_filtros.font = Font(size=9, italic=True, color="52525B")


def _escrever_cabecalho(ws: Worksheet, colunas: Sequence[Coluna]) -> None:
    """Cabeçalho das colunas, com fundo escuro e largura definida."""
    borda = Border(bottom=Side(style="thin", color=_COR_BORDA))
    for indice, coluna in enumerate(colunas, start=1):
        celula = ws.cell(row=_LINHA_CABECALHO, column=indice, value=coluna.titulo)
        celula.font = Font(bold=True, color=_COR_TEXTO_CABECALHO, size=10)
        celula.fill = PatternFill("solid", fgColor=_COR_CABECALHO)
        celula.alignment = Alignment(vertical="center", wrap_text=True)
        celula.border = borda
        ws.column_dimensions[get_column_letter(indice)].width = coluna.largura
    ws.row_dimensions[_LINHA_CABECALHO].height = 24


def _escrever_celula_texto(celula) -> None:
    """
    Força a célula a ser texto, nunca fórmula.

    O openpyxl infere fórmula de qualquer string que comece com ``=`` — então um
    nome, diagnóstico ou observação digitado como ``=HYPERLINK(...)`` viraria
    fórmula viva na planilha exportada (injeção de fórmula, CWE-1236), aberta
    depois na máquina de quem baixou. Fixar ``data_type="s"`` neutraliza isso
    **preservando o texto original**, sem o apóstrofo que o prefixo clássico
    deixaria visível.
    """
    celula.data_type = "s"


def _escrever_linhas(ws: Worksheet, relatorio: Relatorio) -> int:
    """Escreve os dados e devolve o número da última linha preenchida."""
    linha_atual = _LINHA_CABECALHO
    for linha_atual, dados in enumerate(
        relatorio.linhas, start=_LINHA_CABECALHO + 1
    ):
        for indice, coluna in enumerate(relatorio.colunas, start=1):
            valor = _celula_valor(dados.get(coluna.chave))
            celula = ws.cell(row=linha_atual, column=indice, value=valor)
            if isinstance(valor, str):
                _escrever_celula_texto(celula)
            if coluna.formato:
                celula.number_format = coluna.formato
            celula.alignment = Alignment(vertical="top", wrap_text=False)
    return linha_atual


def _escrever_totais(ws: Worksheet, relatorio: Relatorio, ultima_linha: int) -> None:
    """Rodapé com a contagem de registros e a soma das colunas somáveis."""
    linha = ultima_linha + 1
    fundo = PatternFill("solid", fgColor=_COR_TOTAIS)
    borda = Border(top=Side(style="thin", color=_COR_BORDA))

    for indice, coluna in enumerate(relatorio.colunas, start=1):
        celula = ws.cell(row=linha, column=indice)
        celula.fill = fundo
        celula.border = borda
        celula.font = Font(bold=True, size=10)

        if indice == 1:
            celula.value = f"{relatorio.total_registros} registro(s)"
        elif coluna.somar:
            celula.value = _celula_valor(relatorio.totais.get(coluna.chave, 0))
            celula.number_format = coluna.formato or FORMATO_MOEDA


def gerar_planilha(relatorio: Relatorio) -> HttpResponse:
    """
    Monta a planilha do relatório e devolve um ``HttpResponse`` de download.

    O nome do arquivo recebe a data de geração para não sobrescrever exportações
    anteriores na pasta de downloads de quem baixou.
    """
    workbook = Workbook()
    ws = workbook.active
    # O Excel limita o nome da aba a 31 caracteres e proíbe alguns símbolos.
    ws.title = relatorio.titulo[:31]

    _escrever_titulo(ws, relatorio)
    _escrever_cabecalho(ws, relatorio.colunas)
    ultima_linha = _escrever_linhas(ws, relatorio)

    if relatorio.linhas:
        ultima_coluna = get_column_letter(len(relatorio.colunas))
        # Autofiltro e painel congelado: a planilha já chega navegável.
        ws.auto_filter.ref = (
            f"A{_LINHA_CABECALHO}:{ultima_coluna}{ultima_linha}"
        )
        _escrever_totais(ws, relatorio, ultima_linha)
    ws.freeze_panes = f"A{_LINHA_CABECALHO + 1}"

    resposta = HttpResponse(content_type=CONTENT_TYPE_XLSX)
    nome = f"{relatorio.nome_arquivo}_{date.today():%Y-%m-%d}.xlsx"
    resposta["Content-Disposition"] = f'attachment; filename="{nome}"'
    # Sem isso o axios/navegador não enxerga o nome do arquivo em CORS.
    resposta["Access-Control-Expose-Headers"] = "Content-Disposition"
    workbook.save(resposta)
    return resposta


def serializar_colunas(colunas: Iterable[Coluna]) -> list[dict[str, Any]]:
    """Descreve as colunas no JSON, para o frontend montar a tabela de prévia."""
    return [
        {
            "chave": coluna.chave,
            "titulo": coluna.titulo,
            "formato": coluna.formato,
            "somar": coluna.somar,
        }
        for coluna in colunas
    ]
